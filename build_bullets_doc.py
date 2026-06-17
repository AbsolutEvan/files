#!/usr/bin/env python3
"""Bullet-point cue-card version of the Comprehension Debt talk (podium copy)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ALARM = RGBColor(0xC0, 0x2A, 0x14)
GREY  = RGBColor(0x66, 0x66, 0x66)
INK   = RGBColor(0x1A, 0x1A, 0x1A)

doc = Document()
n = doc.styles['Normal']
n.font.name = 'Arial'
n.font.size = Pt(16)
n.paragraph_format.space_after = Pt(6)
n.paragraph_format.line_spacing = 1.12

def cue(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("▶  " + text); r.bold = True; r.font.size = Pt(15)
    r.font.color.rgb = ALARM

def bullet(text, level=0, say=False):
    """level 0/1 bullet. say=True => verbatim punchline (bold, slightly larger)."""
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    if say:
        r.bold = True; r.font.size = Pt(17); r.font.color.rgb = INK
    else:
        r.font.size = Pt(16)
    return p

def stage(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(13); r.font.color.rgb = GREY
    return p

# ---- title ----
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Comprehension Debt"); r.bold = True; r.font.size = Pt(22)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Lightning-talk  ·  ~5 min  ·  bullet cue cards"); r.italic = True
r.font.size = Pt(12); r.font.color.rgb = GREY
leg = doc.add_paragraph()
r = leg.add_run("Bold = say it verbatim.  "); r.bold = True; r.font.size = Pt(11)
r = leg.add_run("Red ▶ = advance the slide.  Grey italic = stage direction."); r.font.size = Pt(11); r.font.color.rgb = GREY

# ===== SLIDE 1 =====
cue("SLIDE 1 up (title)")
bullet("Hi — Evan Wyloge, data editor, Howard Center for Investigative Journalism, ASU")
bullet("My work = where a story collides with a giant, messy pile of data")
bullet("Someone has to vouch for every number in print — that’s this talk")
stage("(pause)")
bullet("First — an old idea from software engineering")

# ===== SLIDE 2 =====
cue("SLIDE 2  (SHIP NOW / PAY LATER)")
bullet("1992 — Ward Cunningham: explaining to his boss why to fix code that already works")
bullet("Reaches for a gut metaphor: debt")
bullet("Ship fast = borrowing. A little is fine — if you pay it back")
bullet("Don’t pay → interest: every future change slower, harder, more painful")
bullet("“Technical debt” — every newsroom has it:")
bullet("the scraper that breaks on every site redesign", level=1)
bullet("the election-night spreadsheet nobody dares touch", level=1)
stage("(slowly)")
bullet("You can feel technical debt. You know you owe it.", say=True)
bullet("Now — a different kind of debt. One you can’t feel.")

# ===== SLIDE 3 =====
cue("SLIDE 3  (df.head(), two header rows, ¯\\_(ツ)_/¯)")
bullet("Earlier this year — helping a student. Big set of public records — not a spreadsheet, millions of rows")
bullet("After the headline numbers: averages + outcomes, by the categories that mattered")
bullet("Moving fast, AI tool doing the analysis. Numbers looked fine — clean, confident, chart-ready")
bullet("I asked them to slow down — show me what it’s actually doing")
bullet("Open a Python script, print first rows → two header rows")
bullet("Two. Why are there two?")
stage("(pause)")
bullet("They didn’t know.", say=True)
bullet("Not a knock — sharp, fast, asked a lot of them. A student I vouch for, who learned from this")
bullet("But sit with it — we’ve all been the person who didn’t know")

# ===== SLIDE 4 =====
cue("SLIDE 4  (terminal: ✨ handled it ✨)")
bullet("Keep pulling the thread:")
bullet("first — delimiter errors", level=1)
bullet("root cause — nul bytes, stray commas: junk buried in millions of rows", level=1)
bullet("Tool hit the junk + never asked how to handle it")
bullet("Made a decision we never saw — dropped the bad rows, mangled the rest into two headers")
stage("(pause)")
bullet("It didn’t fail loudly. It succeeded quietly — on its own terms.", say=True)

# ===== SLIDE 5 =====
cue("SLIDE 5  (green flags → COMPREHENSION DEBT)")
bullet("Looked like technical debt — but we couldn’t feel it. Nothing broke. Every check green")
bullet("It has a name — written about for months: Osmani (Google), O’Reilly, academic studies")
bullet("Comprehension debt = the gap between how much exists and how much anyone understands", say=True)
bullet("Technical debt announces itself. Comprehension debt hides.", say=True)
bullet("Study: AI users finished slightly faster — but scored lower on understanding what they built")
bullet("Fast, confident, and quietly wrong.", say=True)

# ===== SLIDE 6 =====
cue("SLIDE 6  (Which rows got dropped?)")
bullet("The dangerous question isn’t what it found — it’s which rows it dropped")
bullet("Errors aren’t random — they cluster by source / by stretch of time")
bullet("→ dropped rows line up with the exact things we were measuring")
bullet("Numbers quietly biased — no chart would ever show it")
bullet("Remember the scraper you can feel? Our numbers didn’t break, didn’t error")
stage("(pause)")
bullet("They’d simply have been wrong — and we might have published them.", say=True)

# ===== SLIDE 7 =====
cue("SLIDE 7  (the receipt / the rule)")
bullet("What we did: didn’t ban it — slowed it down")
bullet("Rebuilt with AI, carefully → documented script; every decision explicit + ours")
bullet("The move:")
bullet("treat AI output as a draft you can explain back", level=1)
bullet("check what it threw away, not just what it showed you", level=1)
bullet("keep one human who can rebuild the why", level=1)
stage("(pause)")
bullet("Comprehension debt is the technical debt that can get a journalist sued.", say=True)
bullet("The fix isn’t using less AI — it’s never owing more than you can explain.", say=True)
bullet("Thank you.", say=True)

# ===== SLIDE 8 =====
cue("SLIDE 8  (sources) — leave up for Q&A")

# ---- notes ----
doc.add_page_break()
p = doc.add_paragraph(); r = p.add_run("Delivery notes"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = GREY
for txt in [
    "Run time ≈ 5:00–5:25 with the intro. If you’re tight, the cleanest cut is the study sentence on Slide 5 — the story already makes the point.",
    "Two lines to land slow, with silence after: “They didn’t know.” (Slide 3) and “…we might have published them.” (Slide 6).",
    "Before you present: confirm the AI tool’s name (Slides 3–4), and verify the study’s primary source if you name a specific institution out loud.",
]:
    pp = doc.add_paragraph(txt, style='List Bullet')
    for rr in pp.runs: rr.font.size = Pt(12)

doc.save('/home/user/files/Comprehension_Debt_Bullets.docx')
print("saved Comprehension_Debt_Bullets.docx")
