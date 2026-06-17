#!/usr/bin/env python3
"""Build the read-from speaking script for the Comprehension Debt lightning talk."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ALARM = RGBColor(0xC0, 0x2A, 0x14)
GREY  = RGBColor(0x66, 0x66, 0x66)

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Georgia'
st.font.size = Pt(15)          # large for reading at a podium
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.25

def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
    r.font.color.rgb = GREY
    p.paragraph_format.space_before = Pt(14)
    return p

def cue(text):
    p = doc.add_paragraph()
    r = p.add_run("▶  " + text); r.bold = True; r.font.size = Pt(13)
    r.font.color.rgb = ALARM
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    return p

def say(text):
    p = doc.add_paragraph(text)
    return p

def beat(text="(beat — let it sit)"):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(12)
    return p

# ---- Title block ----
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Comprehension Debt"); r.bold = True; r.font.size = Pt(24)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Lightning-talk reading script  ·  ~5 min  ·  read-from copy")
r.italic = True; r.font.color.rgb = GREY; r.font.size = Pt(12)
doc.add_paragraph()

# Delivery legend
leg = doc.add_paragraph()
r = leg.add_run("How to use this: "); r.bold = True; r.font.size = Pt(12)
r = leg.add_run("Read the plain text aloud. ")
r.font.size = Pt(12)
r2 = leg.add_run("Red ▶ lines tell you when to advance the slide"); r2.bold = True; r2.font.color.rgb = ALARM; r2.font.size = Pt(12)
r3 = leg.add_run(" — they are NOT spoken. Italic grey lines are stage directions, also not spoken.")
r3.font.size = Pt(12)

# =================== SLIDE 1 / INTRO ===================
cue("SLIDE 1 up (title) — start here as you walk up")
heading("— Introduce yourself —")
say("Good afternoon. I’m Evan Wyloge, a reporter and editor with the Howard "
    "Center for Investigative Journalism at Arizona State University.")
say("Most of my work lives at the point where a story collides with a giant, "
    "messy pile of data — and somebody has to be able to vouch for every number "
    "that ends up in print. That “vouch for it” part is exactly what this talk "
    "is about.")
beat()
say("The title promises you a tale of two header rows. I’ll get there. But "
    "first, an old idea from software engineering.")

cue("ADVANCE → SLIDE 2  (the “SHIP NOW / PAY LATER” card)")
say("Back in 1992, a programmer named Ward Cunningham had to explain to his "
    "boss why they needed to go back and fix code that already worked. So he "
    "reached for a metaphor the boss would feel in his gut: debt.")
say("When you ship fast, Cunningham said, you’re borrowing. A little debt is "
    "fine — as long as you pay it back. If you don’t, you pay interest: every "
    "future change gets slower, harder, more painful.")
say("They called it technical debt, and every newsroom has it. The scraper "
    "that breaks every time a website is redesigned. The election-night "
    "spreadsheet nobody dares touch. And here’s the important part — ")
beat("slow down")
say("you can feel technical debt. You know you owe it.")
say("I want to tell you about a different kind of debt. One you can’t feel.")

# =================== SLIDE 3 ===================
cue("ADVANCE → SLIDE 3  (df.head(), two header rows, ¯\\_(ツ)_/¯)")
say("Not long ago I was helping someone with a story built on a big set of "
    "public records — not a spreadsheet, millions of rows. We were after the "
    "headline numbers: the key averages and outcomes, broken down by the "
    "categories that mattered.")
say("They were moving fast, with an AI tool doing the analysis. And the "
    "numbers looked great. Clean, confident — the kind you’d drop straight "
    "into a chart and publish.")
say("That’s exactly why I asked them to slow down and show me what the "
    "analysis was actually doing. So we open a Python script, print the first "
    "few rows — and there are two header rows. Two. I asked, why are there two?")
beat("(beat — point at the shrug on screen, then deliver the line flat)")
say("They didn’t know.")
say("That’s not a knock on them — they’re sharp, and they’re fast. But sit "
    "with that answer for a second, because I think every one of us has been "
    "the person who didn’t know.")

# =================== SLIDE 4 ===================
cue("ADVANCE → SLIDE 4  (terminal: “✨ handled it ✨”)")
say("So we keep pulling the thread. First, delimiter errors. Then the root "
    "cause: buried in those millions of rows, a handful of encoding errors — "
    "nul bytes, stray commas. Junk.")
say("And here’s the thing. The tool hit that junk and never asked us how we "
    "wanted to handle it. It just made a decision we never saw — quietly "
    "dropped the rows it didn’t like, and mangled the rest into those two "
    "header rows.")
beat()
say("It didn’t fail loudly. It succeeded quietly — on its own terms.")

# =================== SLIDE 5 ===================
cue("ADVANCE → SLIDE 5  (green checks → COMPREHENSION DEBT)")
say("Now — that looked like technical debt. But we couldn’t feel it. Nothing "
    "broke. Every check was green.")
say("There’s a newer name for this, and engineers have been writing about it "
    "all year — Addy Osmani at Google, O’Reilly, even academic studies. They "
    "call it comprehension debt: the gap between how much exists in your "
    "system and how much any human actually understands.")
say("Technical debt announces itself. Comprehension debt hides. In one "
    "controlled study, people building something with AI finished just as fast "
    "as everyone else — and scored seventeen points lower on understanding "
    "what they’d just built. Fast, confident, and quietly wrong.")

# =================== SLIDE 6 ===================
cue("ADVANCE → SLIDE 6  (“Which rows got dropped?” chart)")
say("Plug that into our world. The dangerous question isn’t what the tool "
    "found. It’s which rows it dropped.")
say("Encoding errors aren’t sprinkled at random. If they cluster around one "
    "source, or one stretch of time, then the rows the tool threw away line up "
    "with the exact things we were measuring. The numbers could be quietly "
    "biased — and no chart would ever show it.")
say("Remember the scraper you can feel? Our numbers didn’t break. They didn’t "
    "throw an error.")
beat()
say("They’d just have been wrong — and we’d have published them.")

# =================== SLIDE 7 ===================
cue("ADVANCE → SLIDE 7  (the receipt / the rule)")
say("So what did we do? We didn’t ban the tool. We slowed it down. We sat "
    "together and worked with an AI tool — carefully — to write a documented "
    "script, where every decision about that junk was explicit, and ours.")
say("That’s the whole move. Treat what the AI hands you as a draft you have to "
    "be able to explain back. Check what it threw away, not just what it showed "
    "you. And keep one human in the room who can rebuild the why.")
beat()
say("Comprehension debt is the version of technical debt that can get a "
    "journalist sued. The fix isn’t using less AI — it’s never owing more than "
    "you can explain.")
say("Thank you.")

# =================== SLIDE 8 ===================
cue("ADVANCE → SLIDE 8  (sources) — leave up for Q&A")
beat("(Nothing to read. Take questions. Sources are on screen.)")

# ---- back-matter notes ----
doc.add_page_break()
heading("Delivery notes")
for n in [
    "Total run time ≈ 5:00–5:25 with the intro. If you’re tight, the cleanest "
    "cut is the “seventeen points lower” sentence on Slide 5 — the story "
    "already proves the point.",
    "The two lines to land slow, with silence after: “They didn’t know.” "
    "(Slide 3) and “…and we’d have published them.” (Slide 6).",
    "Before you present: confirm the AI tool’s name (Slide 3/4 references), and "
    "verify the 17% study’s primary source if you decide to name a specific "
    "institution out loud.",
]:
    p = doc.add_paragraph(n, style='List Bullet');
    for r in p.runs: r.font.size = Pt(12)

doc.save('/home/user/files/Comprehension_Debt_Script.docx')
print("saved Comprehension_Debt_Script.docx")
