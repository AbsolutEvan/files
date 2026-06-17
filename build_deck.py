#!/usr/bin/env python3
"""Build 'The Two Header Rows' lightning-talk deck as a .pptx.
Upload to Google Drive (or File > Import in Slides) to convert to Google Slides."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette ----
PAPER  = RGBColor(0xF7, 0xF5, 0xEF)   # off-white
INK    = RGBColor(0x1A, 0x1A, 0x1A)   # near-black
ALARM  = RGBColor(0xE2, 0x39, 0x1E)   # hot red/orange (only for "silently wrong")
GREY   = RGBColor(0x8A, 0x86, 0x7C)   # fine print
GREEN  = RGBColor(0x2E, 0x7D, 0x32)   # passing checks
TERM_BG= RGBColor(0x1E, 0x1E, 0x1E)   # terminal dark
TERM_FG= RGBColor(0xE6, 0xE6, 0xE6)

SANS = "Arial"          # "you" voice (swap to Inter in Slides if installed)
MONO = "Courier New"    # "the machine" voice

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def box(slide, l, t, w, h, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf

def setline(p, text, font=SANS, size=28, color=INK, bold=False, italic=False,
            align=PP_ALIGN.LEFT):
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    return r

def footer(slide):
    tb, tf = box(slide, SW - Inches(3.2), SH - Inches(0.45),
                 Inches(3.0), Inches(0.35), MSO_ANCHOR.BOTTOM)
    setline(tf.paragraphs[0], "wyloge · howard center",
            MONO, 10, GREY, align=PP_ALIGN.RIGHT)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def newslide(color=PAPER, foot=True):
    s = prs.slides.add_slide(BLANK)
    bg(s, color)
    if foot:
        footer(s)
    return s

# ============ SLIDE 1 — Title (the gag IS the title) ============
s = newslide()
tb, tf = box(s, Inches(1.5), Inches(2.4), Inches(10.3), Inches(2.7))
lines = [
    ("title,subtitle", GREY, 22),
    ("title,subtitle", GREY, 22),
    ("The Two Header Rows,a comprehension debt story", INK, 30),
]
for i, (txt, col, sz) in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    setline(p, txt, MONO, sz, col, bold=(i == 2))
    p.space_after = Pt(6)
notes(s, "(Say nothing about the slide — let them notice the doubled header.) "
         "I want to start with an old idea from software engineering...")

# ============ SLIDE 2 — Technical debt (the credit card) ============
s = newslide()
# card
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(3.67), Inches(2.2), Inches(6.0), Inches(3.0))
card.fill.solid(); card.fill.fore_color.rgb = INK
card.line.color.rgb = INK
ctf = card.text_frame; ctf.word_wrap = True
ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
setline(ctf.paragraphs[0], "SHIP NOW · PAY LATER", SANS, 34, PAPER, bold=True,
        align=PP_ALIGN.CENTER)
p = ctf.add_paragraph(); setline(p, "interest accrues", MONO, 16, ALARM,
                                 align=PP_ALIGN.CENTER)
# footnote
tb, tf = box(s, Inches(3.67), Inches(5.35), Inches(6.0), Inches(0.5),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "the debt metaphor — Ward Cunningham, 1992",
        MONO, 13, GREY, align=PP_ALIGN.CENTER)
notes(s, "Back in 1992, Ward Cunningham needed to explain to his boss why they "
         "had to fix code that already worked — so he reached for a metaphor: "
         "debt. A little debt is fine if you pay it back. The thing about "
         "technical debt is — you can feel it. (Click: the cursed-spreadsheet "
         "image, 'You can feel this one.')")

# ============ SLIDE 3 — "Why are there two?" ============
s = newslide()
tb, tf = box(s, Inches(1.2), Inches(1.3), Inches(8.5), Inches(3.2),
             MSO_ANCHOR.TOP)
code = [
    (">>> df.head()", TERM_FG),
    ("   col_a   col_b   col_c", GREY),
    ("0  col_a   col_b   col_c   <- ???", ALARM),
    ("1  2024..  ...     ...", GREY),
]
# dark code card behind
card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.2),
                          Inches(8.6), Inches(2.6))
card.fill.solid(); card.fill.fore_color.rgb = TERM_BG; card.line.fill.background()
ctf = card.text_frame; ctf.word_wrap = True; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
ctf.margin_left = Inches(0.3)
for i, (txt, col) in enumerate(code):
    p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
    setline(p, txt, MONO, 20, col, bold=(i == 2))
# caption
tb, tf = box(s, Inches(1.0), Inches(4.2), Inches(11.0), Inches(1.6),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "…I don't know.", SANS, 48, INK, bold=True)
notes(s, "We print the first few rows — and there are two header rows. Two. "
         "I asked, why are there two? (beat) They didn't know. "
         "[Screenshot must use FAKE columns, not real data.]")

# ============ SLIDE 4 — "It never asked. It just decided." ============
s = newslide(TERM_BG, foot=False)
tb, tf = box(s, Inches(2.0), Inches(2.6), Inches(9.3), Inches(2.3),
             MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
setline(p, "> found 1,247 malformed rows", MONO, 30, TERM_FG)
p2 = tf.add_paragraph()
r = p2.add_run(); r.text = "> handled it "
r.font.name = MONO; r.font.size = Pt(30); r.font.color.rgb = TERM_FG
r2 = p2.add_run(); r2.text = "✨"
r2.font.name = SANS; r2.font.size = Pt(30); r2.font.color.rgb = ALARM
notes(s, "It hit that junk and never asked how we wanted to handle it. It made "
         "a decision we never saw. It didn't fail loudly. It succeeded quietly "
         "— on its own terms.")

# ============ SLIDE 5 — Name it: comprehension debt ============
s = newslide()
tb, tf = box(s, Inches(2.0), Inches(1.6), Inches(9.3), Inches(3.0),
             MSO_ANCHOR.TOP)
checks = ["✅ TESTS PASS", "✅ ON SCHEDULE", "✅ SHIPPED"]
for i, c in enumerate(checks):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    setline(p, c, SANS, 30, GREEN, bold=True, align=PP_ALIGN.CENTER)
    p.space_after = Pt(10)
# fine print
tb, tf = box(s, Inches(2.0), Inches(4.5), Inches(9.3), Inches(0.5),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "(nobody understands this)", SANS, 14, GREY,
        italic=True, align=PP_ALIGN.CENTER)
# the term (reveal)
tb, tf = box(s, Inches(1.5), Inches(5.2), Inches(10.3), Inches(1.7),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "COMPREHENSION DEBT", SANS, 40, ALARM, bold=True,
        align=PP_ALIGN.CENTER)
p = tf.add_paragraph()
setline(p, "the gap between how much exists and how much anyone understands",
        SANS, 16, INK, italic=True, align=PP_ALIGN.CENTER)
notes(s, "That looked like technical debt. But we couldn't feel it — nothing "
         "broke. Engineers have a name for it: comprehension debt (Osmani, "
         "O'Reilly, academic studies). Technical debt announces itself. "
         "Comprehension debt hides. One study: just as fast, 17 points lower on "
         "understanding what they built.")

# ============ SLIDE 6 — Which rows got dropped? ============
s = newslide()
tb, tf = box(s, Inches(0.8), Inches(0.6), Inches(11.0), Inches(0.9),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "Which rows got dropped?", SANS, 34, INK, bold=True)
# bars
bar_vals = [2.6, 3.4, 2.0, 3.0]
ghost = [False, False, True, True]  # last two faded
x0 = Inches(2.0); base = Inches(5.8); bw = Inches(1.6); gap = Inches(0.6)
for i, (v, g) in enumerate(zip(bar_vals, ghost)):
    h = Inches(v)
    x = x0 + i * (bw + gap)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, base - h, bw, h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0xCF, 0xCB, 0xC0) if g else INK
    bar.line.fill.background()
    lbl = s.shapes.add_textbox(x, base + Inches(0.05), bw, Inches(0.4))
    setline(lbl.text_frame.paragraphs[0], chr(ord('A') + i), MONO, 16, GREY,
            align=PP_ALIGN.CENTER)
# trash-can label
tb, tf = box(s, Inches(9.6), Inches(2.0), Inches(3.0), Inches(1.0),
             MSO_ANCHOR.MIDDLE)
setline(tf.paragraphs[0], "\U0001f5d1 probably nothing", SANS, 18, ALARM, bold=True,
        align=PP_ALIGN.CENTER)
notes(s, "The dangerous question isn't what it found — it's which rows it "
         "dropped. Errors aren't random; the dropped rows can line up with the "
         "exact things you're measuring. Remember the scraper you can feel? Our "
         "numbers didn't break. They'd just have been wrong — and we'd have "
         "published them. [Use neutral labels, not real data.]")

# ============ SLIDE 7 — The rule (receipt) ============
s = newslide()
# receipt rectangle
rec = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.17), Inches(1.2),
                         Inches(5.0), Inches(4.0))
rec.fill.solid(); rec.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
rec.line.color.rgb = GREY
rtf = rec.text_frame; rtf.word_wrap = True; rtf.vertical_anchor = MSO_ANCHOR.TOP
rtf.margin_left = Inches(0.35); rtf.margin_top = Inches(0.3)
items = [
    ("RECEIPT", 16, INK, True),
    ("analysis ............... done", 13, GREY, False),
    ("cleanup ................ done", 13, GREY, False),
    ("understanding .......... $0.00", 13, ALARM, False),
    ("", 8, GREY, False),
    ("------------------------------", 13, GREY, False),
]
for i, (txt, sz, col, b) in enumerate(items):
    p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
    setline(p, txt, MONO, sz, col, bold=b, align=PP_ALIGN.CENTER)
# stamp
stamp, stf = box(s, Inches(2.2), Inches(3.0), Inches(9.0), Inches(1.6),
                 MSO_ANCHOR.MIDDLE)
setline(stf.paragraphs[0], "NEVER OWE MORE\nTHAN YOU CAN EXPLAIN", SANS, 36, ALARM,
        bold=True, align=PP_ALIGN.CENTER)
# rotate stamp for rubber-stamp feel
stamp.rotation = -8
notes(s, "We didn't ban the tool. We slowed it down and wrote a documented "
         "script where every decision was explicit, and ours. The fix isn't "
         "using less AI — it's never owing more than you can explain. Thank you.")

# ============ SLIDE 8 — Sources (leave-up for Q&A) ============
s = newslide()
tb, tf = box(s, Inches(1.0), Inches(0.8), Inches(11.3), Inches(0.9),
             MSO_ANCHOR.TOP)
setline(tf.paragraphs[0], "yes, this is a real field of study now", SANS, 24,
        INK, italic=True)
tb, tf = box(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.5),
             MSO_ANCHOR.TOP)
refs = [
    'Cunningham, W. (1992) — origin of "technical debt"',
    "Osmani, A. (2026) — Comprehension Debt: The Hidden Cost of AI-Generated Code",
    "O'Reilly Radar — same essay",
    "arXiv 2512.08942 — Beyond Technical Debt (study)",
    "Wang, S. — Managing Comprehension Debt (ITNext)",
]
for i, r in enumerate(refs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    setline(p, r, MONO, 16, INK)
    p.space_after = Pt(10)
notes(s, "(Q&A leave-up slide. Optional: add a QR code to a links doc.)")

prs.save("/home/user/files/The_Two_Header_Rows.pptx")
print("saved The_Two_Header_Rows.pptx with", len(prs.slides.slides_elements) if hasattr(prs.slides,'slides_elements') else len(prs.slides._sldIdLst), "slides")
